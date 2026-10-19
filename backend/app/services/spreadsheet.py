"""Чтение таблиц для импорта и сборка шаблонов.

Принимаются .xlsx и .csv. Заголовки сравниваются без регистра и лишних пробелов,
«ё» считается за «е», поэтому «Название», «название » и «НАЗВАНИЕ» — одна колонка.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class TableError(ValueError):
    """Файл не прочитать или в нём нет обязательной колонки."""


def normalize(value: object) -> str:
    return " ".join(str(value or "").replace("ё", "е").replace("Ё", "Е").split()).lower()


@dataclass(frozen=True)
class Column:
    key: str
    title: str  # как колонка называется в шаблоне
    aliases: tuple[str, ...] = ()
    required: bool = False
    note: str = ""

    def matches(self, header: str) -> bool:
        return header in {normalize(self.title), *(normalize(a) for a in self.aliases)}


def _rows_from_xlsx(content: bytes) -> list[list[object]]:
    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception as exc:  # openpyxl бросает разное на битых файлах
        raise TableError("Не удалось открыть файл Excel. Сохраните его как .xlsx") from exc
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _rows_from_csv(content: bytes) -> list[list[object]]:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise TableError("Не удалось прочитать кодировку файла")
    delimiter = ";" if text.split("\n", 1)[0].count(";") >= text.split("\n", 1)[0].count(",") \
        else ","
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_table(content: bytes, filename: str, columns: list[Column]) -> list[tuple[int, dict]]:
    """Строки таблицы как (номер строки в файле, {ключ колонки: значение}).

    Пустые строки пропускаются, лишние колонки не читаются.
    """
    if not content:
        raise TableError("Файл пустой")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        rows = _rows_from_xlsx(content)
    elif name.endswith((".csv", ".txt")):
        rows = _rows_from_csv(content)
    else:
        raise TableError("Нужен файл .xlsx или .csv")
    # Заголовок — первая непустая строка
    start = next((i for i, row in enumerate(rows) if any(normalize(c) for c in row)), None)
    if start is None:
        raise TableError("В файле нет данных")
    headers = [normalize(c) for c in rows[start]]
    positions: dict[str, int] = {}
    for column in columns:
        index = next((i for i, h in enumerate(headers) if column.matches(h)), None)
        if index is not None:
            positions[column.key] = index
        elif column.required:
            raise TableError(f"Нет обязательной колонки «{column.title}»")

    result = []
    for number, row in enumerate(rows[start + 1:], start=start + 2):
        values = {}
        for key, index in positions.items():
            value = row[index] if index < len(row) else None
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            values[key] = "" if value is None else str(value).strip()
        if any(values.values()):
            result.append((number, values))
    return result


def number(value: str) -> float | None:
    """Число из ячейки: десятичная запятая допускается. Пусто — None."""
    value = (value or "").replace(" ", "").replace(",", ".")
    if not value:
        return None
    return float(value)


def build_template(columns: list[Column], examples: list[list[str]], help_rows: list[str],
                   sheet_title: str) -> bytes:
    """Шаблон: лист с заголовками и примерами и лист «Как заполнять»."""
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_title
    head_fill = PatternFill("solid", fgColor="E9F1FC")
    required_fill = PatternFill("solid", fgColor="FFE7C2")
    for i, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=i, value=column.title)
        cell.font = Font(bold=True)
        cell.fill = required_fill if column.required else head_fill
        sheet.column_dimensions[get_column_letter(i)].width = max(len(column.title) + 6, 18)
    for r, example in enumerate(examples, start=2):
        for c, value in enumerate(example, start=1):
            sheet.cell(row=r, column=c, value=value)
    sheet.freeze_panes = "A2"

    notes = book.create_sheet("Как заполнять")
    notes.column_dimensions["A"].width = 28
    notes.column_dimensions["B"].width = 90
    notes.cell(row=1, column=1, value="Колонка").font = Font(bold=True)
    notes.cell(row=1, column=2, value="Что писать").font = Font(bold=True)
    for r, column in enumerate(columns, start=2):
        notes.cell(row=r, column=1, value=column.title + (" *" if column.required else ""))
        notes.cell(row=r, column=2, value=column.note).alignment = Alignment(wrap_text=True)
    row = len(columns) + 3
    notes.cell(row=row, column=1, value="* — обязательная колонка").font = Font(italic=True)
    for text in help_rows:
        row += 1
        notes.cell(row=row, column=1, value=text)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
