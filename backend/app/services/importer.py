"""Импорт штатных расписаний и контингента из CSV/XLSX.

Формат подогнан под типовые выгрузки: названия колонок могут быть русскими
(как в отчётах 1С:ЗУП) или английскими. Персональные данные отсекаются здесь же:
ФИО или табельный номер превращаются в хэш, а колонки, которых нет в списке
известных (паспорт, СНИЛС, адрес и т. п.), просто не читаются.
"""

from __future__ import annotations

import csv
import hashlib
import hmac
import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import Enrollment, Level, Organization, StaffPosition, StudyForm
from app.services.functions import BY_CODE, classify_title

STAFF_COLUMNS = {
    "organization": ("organization", "организация", "учреждение", "оу", "код оу"),
    "department": ("department", "подразделение", "структурное подразделение"),
    "title": ("title", "должность", "наименование должности"),
    "rate": ("rate", "ставка", "ставки", "количество ставок", "кол-во ставок",
             "количество штатных единиц"),
    "monthly_pay": ("monthly_pay", "фот", "начислено", "фот в месяц", "месячный фонд"),
    "employee": ("employee", "сотрудник", "фио", "работник", "табельный номер"),
    "weekly_hours": ("weekly_hours", "часов в неделю", "нагрузка", "недельная нагрузка"),
    "duties": ("duties", "функции", "обязанности", "должностные обязанности"),
    "function_code": ("function_code", "код функции"),
}

ENROLLMENT_COLUMNS = {
    "organization": STAFF_COLUMNS["organization"],
    "academic_year": ("academic_year", "учебный год", "год"),
    "level": ("level", "ступень", "уровень"),
    "study_form": ("study_form", "форма обучения", "форма"),
    "classes_count": ("classes_count", "классов", "классы", "групп", "классов-комплектов"),
    "students_count": ("students_count", "обучающихся", "учащихся", "численность", "контингент"),
    "special_needs_count": ("special_needs_count", "овз", "с овз", "детей с овз"),
}

LEVEL_ALIASES = {
    "preschool": Level.preschool, "дошкольное": Level.preschool, "доу": Level.preschool,
    "primary": Level.primary, "начальное": Level.primary, "1-4": Level.primary,
    "basic": Level.basic, "основное": Level.basic, "5-9": Level.basic,
    "secondary": Level.secondary, "среднее": Level.secondary, "10-11": Level.secondary,
}
FORM_ALIASES = {
    "full_time": StudyForm.full_time, "очная": StudyForm.full_time, "": StudyForm.full_time,
    "distance": StudyForm.distance, "дистанционная": StudyForm.distance,
    "home": StudyForm.home, "на дому": StudyForm.home, "надомная": StudyForm.home,
}


class ImportFormatError(ValueError):
    """Файл нельзя загрузить целиком (нет обязательных колонок, не тот формат)."""


@dataclass
class ImportResult:
    imported: int = 0
    skipped: int = 0
    organizations: int = 0
    dropped_columns: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    def error(self, row_number: int, message: str) -> None:
        self.skipped += 1
        if len(self.errors) < 50:
            self.errors.append(f"строка {row_number}: {message}")


def pseudonymize(value: str) -> str:
    """Необратимый идентификатор сотрудника.

    HMAC с секретной солью, а не просто SHA-256: без соли ФИО легко
    восстановить перебором по списку сотрудников.
    """
    normalized = " ".join(value.lower().replace("ё", "е").split())
    salt = get_settings().pii_hash_salt.encode()
    return hmac.new(salt, normalized.encode(), hashlib.sha256).hexdigest()[:32]


def read_table(filename: str, content: bytes) -> list[dict[str, str]]:
    """Читает CSV или XLSX в список словарей с нормализованными заголовками."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        rows = [["" if c is None else str(c).strip() for c in row]
                for row in sheet.iter_rows(values_only=True)]
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1251")  # так по умолчанию сохраняет Excel
        delimiter = ";" if text.count(";") >= text.count(",") else ","
        rows = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text),
                                                               delimiter=delimiter)]
    rows = [row for row in rows if any(row)]
    if len(rows) < 2:
        raise ImportFormatError("В файле нет строк с данными")
    header = [h.lower().replace("ё", "е") for h in rows[0]]
    return [dict(zip(header, row, strict=False)) for row in rows[1:]]


def _map_columns(header: list[str], aliases: dict[str, tuple[str, ...]], required: tuple[str, ...],
                 result: ImportResult) -> dict[str, str]:
    mapping = {}
    known = set()
    for field_name, names in aliases.items():
        for column in header:
            if column in names:
                mapping[field_name] = column
                known.add(column)
                break
    missing = [aliases[name][1] for name in required if name not in mapping]
    if missing:
        raise ImportFormatError("Не найдены обязательные колонки: " + ", ".join(missing))
    result.dropped_columns = [c for c in header if c and c not in known]
    return mapping


def _number(value: str, default: str | None = None) -> Decimal:
    cleaned = (value or "").replace(" ", "").replace(" ", "").replace(",", ".")
    if not cleaned and default is not None:
        cleaned = default
    try:
        return Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"«{value}» не число") from exc


def _find_organizations(db: Session, allowed_ids: set[int] | None) -> dict[str, Organization]:
    index = {}
    for org in db.scalars(select(Organization)):
        if allowed_ids is not None and org.id not in allowed_ids:
            continue
        index[org.name.lower()] = org
        index[str(org.id)] = org
        if org.external_id:
            index[org.external_id.lower()] = org
    return index


def import_staffing(db: Session, rows: list[dict[str, str]], *, allowed_org_ids: set[int] | None,
                    replace: bool = True) -> ImportResult:
    """Загружает штатное расписание. replace=True заменяет расписание тех ОУ,
    которые встретились в файле, — обычный сценарий «загрузили свежую выгрузку»."""
    result = ImportResult()
    columns = _map_columns(list(rows[0]), STAFF_COLUMNS, ("organization", "title", "rate"), result)
    organizations = _find_organizations(db, allowed_org_ids)

    positions: list[StaffPosition] = []
    touched: set[int] = set()
    for number, row in enumerate(rows, start=2):
        def cell(name: str, row=row) -> str:
            return row.get(columns.get(name, ""), "") or ""

        org = organizations.get(cell("organization").lower())
        if org is None:
            result.error(number, f"учреждение «{cell('organization')}» не найдено или недоступно")
            continue
        title = cell("title")
        if not title:
            result.error(number, "не указана должность")
            continue
        try:
            rate = _number(cell("rate"))
            pay = _number(cell("monthly_pay"), "0")
            hours = _number(cell("weekly_hours")) if cell("weekly_hours") else None
        except ValueError as exc:
            result.error(number, str(exc))
            continue
        if not Decimal("0") < rate <= Decimal("3"):
            result.error(number, f"ставка {rate} вне допустимого диапазона")
            continue

        function = BY_CODE.get(cell("function_code")) or classify_title(title)
        employee = cell("employee")
        positions.append(StaffPosition(
            organization_id=org.id,
            department=cell("department") or None,
            title=title[:200],
            category=function.category.value,
            function_code=function.code,
            rate=rate,
            monthly_pay=pay,
            employee_ref=pseudonymize(employee) if employee else None,
            weekly_hours=hours,
            duties=cell("duties") or None,
        ))
        touched.add(org.id)

    if replace and touched:
        db.execute(delete(StaffPosition).where(StaffPosition.organization_id.in_(touched)))
    db.add_all(positions)
    result.imported = len(positions)
    result.organizations = len(touched)
    return result


def import_enrollment(db: Session, rows: list[dict[str, str]], *,
                      allowed_org_ids: set[int] | None) -> ImportResult:
    result = ImportResult()
    columns = _map_columns(
        list(rows[0]), ENROLLMENT_COLUMNS,
        ("organization", "academic_year", "level", "students_count"), result,
    )
    organizations = _find_organizations(db, allowed_org_ids)
    touched: set[int] = set()

    for number, row in enumerate(rows, start=2):
        def cell(name: str, row=row) -> str:
            return row.get(columns.get(name, ""), "") or ""

        org = organizations.get(cell("organization").lower())
        level = LEVEL_ALIASES.get(cell("level").lower())
        form = FORM_ALIASES.get(cell("study_form").lower())
        if org is None:
            result.error(number, f"учреждение «{cell('organization')}» не найдено или недоступно")
            continue
        if level is None or form is None:
            result.error(number, "не распознана ступень или форма обучения")
            continue
        try:
            # «2026/2027» и «2026-27» сводим к году начала
            year = int(cell("academic_year")[:4])
            values = {
                "classes_count": int(_number(cell("classes_count"), "0")),
                "students_count": int(_number(cell("students_count"))),
                "special_needs_count": int(_number(cell("special_needs_count"), "0")),
            }
        except ValueError as exc:
            result.error(number, str(exc))
            continue

        existing = db.scalar(select(Enrollment).where(
            Enrollment.organization_id == org.id, Enrollment.academic_year == year,
            Enrollment.level == level.value, Enrollment.study_form == form.value,
        ))
        if existing is None:
            existing = Enrollment(organization_id=org.id, academic_year=year,
                                  level=level.value, study_form=form.value)
            db.add(existing)
        for key, value in values.items():
            setattr(existing, key, value)
        db.flush()
        touched.add(org.id)
        result.imported += 1

    result.organizations = len(touched)
    return result
