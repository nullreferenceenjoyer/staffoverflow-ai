"""Бригады и сотрудники на точке, трудоёмкость, этапы, учёт учреждений и импорт."""

import io

from openpyxl import Workbook, load_workbook

ADMIN = ("admin@example.ru", "admin")
DIRECTOR = "director@demo.ru"
MUNICIPAL = "municipal@demo.ru"
EXECUTOR = "executor@demo.ru"


def admin(login):
    return login(*ADMIN)


def director_org(client, login) -> int:
    return client.get("/api/auth/me", headers=login(DIRECTOR)).json()["organization_id"]


def brigade_of(client, login, org_id: int, topic: str = "plumbing") -> dict:
    """Общая бригада (без специализации или с этой темой), за которой закреплено учреждение."""
    brigades = client.get("/api/admin/brigades", headers=admin(login)).json()
    return next(b for b in brigades
                if any(o["id"] == org_id for o in b["organizations"])
                and (not b["topics"] or topic in b["topics"]))


def xlsx(rows: list[list]) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_simple_work_goes_to_staff_on_site(client, login):
    body = client.post("/api/requests", headers=login(DIRECTOR),
                       json={"title": "Отремонтировать дверцу шкафа в кабинете 3"}).json()
    assert body["complexity"] == "simple"
    assert body["assignee"]["role"] == "executor"
    assert body["brigade"] is None


def test_complex_work_goes_to_brigade_and_any_member_takes_it(client, login):
    org_id = director_org(client, login)
    brigade = brigade_of(client, login, org_id)
    body = client.post("/api/requests", headers=login(DIRECTOR), json={
        "title": "Заменить трубы отопления в спортзале"}).json()
    assert body["complexity"] == "complex"
    assert body["effort_hours"] >= 6
    assert body["brigade"]["id"] == brigade["id"]
    assert (body["status"], body["assignee"]) == ("assigned", None)
    # Срок длиннее типового для сантехники (3 дня): работа на полтора дня
    assert body["due_at"][:10] > body["created_at"][:10]

    member = login(brigade["members"][0]["email"], "demo12345")
    seen = client.get(f"/api/requests/{body['id']}", headers=member).json()
    assert seen["my_turn"].startswith("Возьмите заявку в работу")
    # Сотрудник на точке её не видит: она не для него
    assert client.get(f"/api/requests/{body['id']}", headers=login(EXECUTOR)).status_code == 404

    taken = client.post(f"/api/requests/{body['id']}/status", headers=member,
                        json={"status": "in_progress"}).json()
    assert taken["status"] == "in_progress"
    assert taken["assignee"]["name"] == brigade["members"][0]["name"]
    assert taken["brigade"]["id"] == brigade["id"]


def test_major_work_is_flagged_and_can_be_split_into_stages(client, login):
    municipal = login(MUNICIPAL)
    body = client.post("/api/requests", headers=login(DIRECTOR), json={
        "title": "Покрасить стены во всех кабинетах третьего этажа"}).json()
    assert body["complexity"] == "major"
    # Автор заявки — директор, поэтому предупреждение получает специалист округа
    notes = client.get("/api/notifications", headers=municipal).json()
    assert any(n["request_id"] == body["id"] and "крупная работа" in n["text"] for n in notes)

    url = f"/api/requests/{body['id']}"
    assert client.post(f"{url}/split", headers=municipal,
                       json={"stages": [{"title": "Один этап"}]}).status_code == 422
    split = client.post(f"{url}/split", headers=municipal, json={"stages": [
        {"title": "Подготовить стены", "hours": 8}, {"title": "Покрасить", "hours": 10}]})
    assert split.status_code == 201, split.text
    parent = split.json()
    assert parent["status"] == "in_progress"
    assert [c["title"] for c in parent["children"]] == ["Подготовить стены", "Покрасить"]
    assert client.post(f"{url}/split", headers=municipal, json={"stages": [
        {"title": "Ещё раз"}, {"title": "И ещё"}]}).status_code == 409

    # Этапы выполнены — общая заявка ждёт приёмки у заявителя
    for child in parent["children"]:
        stage = client.get(f"/api/requests/{child['id']}", headers=municipal).json()
        assert stage["parent"]["id"] == body["id"]
        assert stage["effort_hours"] in (8, 10)
        done = client.post(f"/api/requests/{child['id']}/status", headers=municipal,
                           json={"status": "done", "comment": "Сделано"})
        assert done.status_code == 200, done.text
    after = client.get(url, headers=login(DIRECTOR)).json()
    assert after["status"] == "done"
    assert after["my_turn"] == "Проверьте работу и примите её"


def test_manager_corrects_estimate_and_hands_over_to_brigade(client, login):
    director = login(DIRECTOR)
    body = client.post("/api/requests", headers=director,
                       json={"title": "Починить стул в учительской"}).json()
    assert body["assignee"]["role"] == "executor"
    url = f"/api/requests/{body['id']}"
    corrected = client.patch(url, headers=director, json={"effort_hours": 9}).json()
    assert (corrected["complexity"], corrected["effort_by"]) == ("complex", "manual")
    assert corrected["due_at"] > body["due_at"]

    # Передать бригаде может только управление образования
    municipal = login(MUNICIPAL)
    details = client.get(url, headers=municipal).json()
    serving = [b for b in details["brigades"] if b["serves"]]
    assert serving
    assert client.patch(url, headers=director,
                        json={"brigade_id": serving[0]["id"]}).status_code == 403
    moved = client.patch(url, headers=municipal, json={"brigade_id": serving[0]["id"]}).json()
    assert moved["brigade"]["id"] == serving[0]["id"]
    assert moved["assignee"] is None
    assert client.patch(url, headers=municipal, json={"brigade_id": 999_999}).status_code == 422


def test_organization_without_staff_on_site_is_served_by_brigade(client, login):
    headers = admin(login)
    orgs = client.get("/api/admin/organizations", headers=headers).json()
    # В демо-данных хозслужба есть только в учреждениях пилотного округа
    without_staff = next(o for o in orgs if o["brigades"] and o["users"] == 0)
    body = client.post("/api/requests", headers=headers, json={
        "title": "Отремонтировать дверцу шкафа", "organization_id": without_staff["id"]}).json()
    assert body["complexity"] == "simple"
    assert body["brigade"]["id"] in {b["id"] for b in without_staff["brigades"]}


def test_admin_keeps_register_of_organizations(client, login):
    headers = admin(login)
    created = client.post("/api/admin/organizations", headers=headers, json={
        "name": "Гимназия № 77", "municipality": "Новый округ", "org_type": "school",
        "address": "г. Новый, ул. Мира, д. 1", "latitude": 55.5, "longitude": 37.5,
        "buildings": [{"name": "Главный корпус"}, {"name": "Спортзал"}]})
    assert created.status_code == 201, created.text
    org = created.json()
    assert org["municipality"] == "Новый округ"
    assert [b["name"] for b in org["buildings"]] == ["Главный корпус", "Спортзал"]

    keep = org["buildings"][0]
    updated = client.put(f"/api/admin/organizations/{org['id']}", headers=headers, json={
        "name": "Гимназия № 77 имени Пушкина", "municipality": "Новый округ",
        "org_type": "school", "buildings": [{**keep, "name": "Корпус А"}, {"name": "Корпус Б"}]})
    assert [b["name"] for b in updated.json()["buildings"]] == ["Корпус А", "Корпус Б"]

    assert client.get("/api/admin/organizations", headers=login(DIRECTOR)).status_code == 403
    used = next(o for o in client.get("/api/admin/organizations", headers=headers).json()
                if o["users"] > 0)
    assert client.delete(f"/api/admin/organizations/{used['id']}",
                         headers=headers).status_code == 409
    assert client.delete(f"/api/admin/organizations/{org['id']}",
                         headers=headers).status_code == 204


def test_templates_are_valid_excel_files(client, login):
    for kind, first in (("organizations", "Название"), ("users", "ФИО")):
        response = client.get(f"/api/admin/templates/{kind}.xlsx", headers=admin(login))
        assert response.status_code == 200
        book = load_workbook(io.BytesIO(response.content))
        assert book.worksheets[0]["A1"].value == first
        assert book.worksheets[1].title == "Как заполнять"
    assert client.get("/api/admin/templates/other.xlsx", headers=admin(login)).status_code == 404


def test_import_organizations_from_excel(client, login):
    headers = admin(login)
    content = xlsx([
        ["Название", "Муниципалитет", "Адрес", "Широта", "Долгота", "Корпуса", "Код"],
        ["СОШ № 101", "Импортный округ", "ул. Первая, 1", "55,61", "37,42",
         "Главный корпус; Мастерские", "IMP-101"],
        ["Детский сад № 102", "", "", "", "", "", ""],
        ["", "Импортный округ", "", "", "", "", ""],
        ["Школа без координат", "Импортный округ", "", "abc", "", "", ""],
    ])
    response = client.post("/api/admin/import/organizations", headers=headers,
                           files={"file": ("orgs.xlsx", content)},
                           data={"municipality": "Импортный округ"})
    assert response.status_code == 200, response.text
    result = response.json()
    assert (result["created"], result["updated"], result["skipped"]) == (2, 0, 2)
    assert result["created_municipalities"] == ["Импортный округ"]
    assert [e["row"] for e in result["errors"]] == [4, 5]

    orgs = {o["name"]: o for o in client.get("/api/admin/organizations", headers=headers).json()}
    school = orgs["СОШ № 101"]
    assert [b["name"] for b in school["buildings"]] == ["Главный корпус", "Мастерские"]
    assert (school["latitude"], school["org_type"]) == (55.61, "school")
    assert orgs["Детский сад № 102"]["org_type"] == "preschool"
    assert [b["name"] for b in orgs["Детский сад № 102"]["buildings"]] == ["Основное здание"]

    # Повторная загрузка по коду обновляет, а не создаёт дубль
    again = xlsx([["Наименование", "Код"], ["СОШ № 101 им. Гагарина", "IMP-101"]])
    result = client.post("/api/admin/import/organizations", headers=headers,
                         files={"file": ("orgs.xlsx", again)}).json()
    assert (result["created"], result["updated"]) == (0, 1)

    no_name = xlsx([["Адрес"], ["ул. Вторая"]])
    response = client.post("/api/admin/import/organizations", headers=headers,
                           files={"file": ("orgs.xlsx", no_name)})
    assert response.status_code == 422
    assert "Название" in response.json()["detail"]


def test_import_users_creates_passwords_and_binds_them(client, login):
    headers = admin(login)
    org = next(o for o in client.get("/api/admin/organizations", headers=headers).json()
               if o["users"] > 0)
    brigade = client.get("/api/admin/brigades", headers=headers).json()[0]
    content = xlsx([
        ["ФИО", "E-mail", "Роль", "Учреждение", "Муниципалитет", "Бригада", "Пароль"],
        ["Новикова Т. С.", "novikova@demo.ru", "Сотрудник", org["name"], "", "", ""],
        ["Белкин Р. А.", "belkin@demo.ru", "Выездная бригада", "", "", brigade["name"],
         "brigada2026"],
        ["Ошибкин Н. Н.", "oshibkin@demo.ru", "Директор", "Несуществующая школа", "", "", ""],
        ["Кто-то", "не-почта", "Сотрудник", org["name"], "", "", ""],
        ["Лебедева О. И.", "employee@demo.ru", "сотрудник", org["name"], "", "", ""],
    ])
    result = client.post("/api/admin/import/users", headers=headers,
                         files={"file": ("users.xlsx", content)}).json()
    assert (result["created"], result["updated"], result["skipped"]) == (2, 1, 2)
    assert [p["email"] for p in result["passwords"]] == ["novikova@demo.ru"]
    assert {e["row"] for e in result["errors"]} == {4, 5}

    password = result["passwords"][0]["password"]
    me = client.get("/api/auth/me", headers=login("novikova@demo.ru", password)).json()
    assert (me["role"], me["organization_id"]) == ("employee", org["id"])
    member = client.get("/api/auth/me", headers=login("belkin@demo.ru", "brigada2026")).json()
    assert member["brigade_id"] == brigade["id"]


def test_brigade_schools_are_suggested_by_distance(client, login):
    headers = admin(login)
    municipalities = client.get("/api/municipalities", headers=headers).json()
    pilot = municipalities[0]
    member = client.post("/api/users", headers=headers, json={
        "email": "new.brigade@demo.ru", "full_name": "Сергеев О. П.", "password": "secret123",
        "role": "brigade"})
    assert member.status_code == 422  # без бригады участника не завести

    created = client.post("/api/admin/brigades", headers=headers, json={
        "name": "Бригада «Центр»", "municipality_id": pilot["id"],
        "latitude": 55.9, "longitude": 37.0, "topics": ["plumbing"]})
    assert created.status_code == 201, created.text
    brigade = created.json()
    assert client.post("/api/admin/brigades", headers=headers, json={
        "name": "Бригада «Центр»"}).status_code == 409
    assert client.post("/api/admin/brigades", headers=headers, json={
        "name": "Неизвестная", "topics": ["norms"]}).status_code == 422

    user = client.post("/api/users", headers=headers, json={
        "email": "new.brigade@demo.ru", "full_name": "Сергеев О. П.", "password": "secret123",
        "role": "brigade", "brigade_id": brigade["id"]}).json()
    updated = client.put(f"/api/admin/brigades/{brigade['id']}", headers=headers, json={
        "name": "Бригада «Центр»", "municipality_id": pilot["id"], "latitude": 55.9,
        "longitude": 37.0, "topics": ["plumbing"], "member_ids": [user["id"]]}).json()
    assert [m["id"] for m in updated["members"]] == [user["id"]]

    suggestions = client.get(f"/api/admin/brigades/{brigade['id']}/suggestions",
                             headers=headers).json()
    assert suggestions["has_base"]
    distances = [i["distance_km"] for i in suggestions["items"] if i["distance_km"] is not None]
    assert distances == sorted(distances)
    assert all(i["nearest"] for i in suggestions["items"] if i["distance_km"] is not None)

    chosen = [i["id"] for i in suggestions["items"][:2]]
    saved = client.put(f"/api/admin/brigades/{brigade['id']}/organizations", headers=headers,
                       json={"organization_ids": chosen}).json()
    assert {o["id"] for o in saved["organizations"]} == set(chosen)
    assert all(o["distance_km"] is not None for o in saved["organizations"])
